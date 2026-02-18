"""
Article Generator - Core Logic (Web Version)
Extracted from the original script and adapted for use as a web application.
- Accepts a progress callback instead of printing directly
- Accepts an output directory for saving files
- No hardcoded values or input() calls
"""
import pandas as pd
import anthropic
import os
import re
import time
import requests
from pathlib import Path
from typing import Optional, Callable

try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


DEFAULT_ARTICLE_INSTRUCTIONS = """Write an SEO optimized article of 2000-2500 words using everyday English. Include naturally fitting keywords ({KEYWORDS}) in the article. Write from first person plural perspective (we, us, or using the brand name "{CLIENT}").
Avoid using AI words like dive, delve, tapestry, etc.
Do not include any fake examples, reviews, etc.
Content Requirements:
- Follow the article outline/instructions provided above closely
- Use small paragraphs for readability
- Include one numbered list, one bullet list and two tables
- Base content on the company background provided above
End the article with a strong concluding section that also acts as a CTA (maximum 3 short paragraphs) that also serves as a CTA for the client. Use an appropriate bold anchor text (up to 4 words) to link back to {LINKS}. Give a strong heading to this section.
If pricing is involved, add a small one-line disclaimer at the end of the article and format it in italics using *italic text* format.
Write the complete article in markdown format with:
- **bold text** for keywords and emphasis
- *italic text* for disclaimers
- [link text](URL) for hyperlinks
- Proper table formatting with | pipes |
- Lists using - or numbered format"""


class ArticleGeneratorPureDocx:
    def __init__(self, api_key: str, progress_callback: Optional[Callable] = None):
        """Initialize with Anthropic API key and optional progress callback."""
        self.client = anthropic.Anthropic(api_key=api_key)
        self._callback = progress_callback

    def log(self, message: str):
        """Send a progress message via callback or print as fallback."""
        if self._callback:
            self._callback(message)
        else:
            print(message)

    def sanitize_filename(self, title: str) -> str:
        filename = re.sub(r'[<>:"/\\|?*]', '', title)
        filename = re.sub(r'\s+', ' ', filename).strip()
        if len(filename) > 200:
            filename = filename[:200].rsplit(' ', 1)[0]
        return filename

    def sanitize_folder_name(self, client_name: str) -> str:
        folder_name = re.sub(r'[<>:"/\\|?*]', '', client_name)
        folder_name = re.sub(r'\s+', ' ', folder_name).strip()
        if len(folder_name) > 100:
            folder_name = folder_name[:100].rsplit(' ', 1)[0]
        if not folder_name:
            folder_name = "Unknown_Client"
        return folder_name

    def fetch_website_content(self, url: str) -> str:
        if not BS4_AVAILABLE:
            self.log("   ⚠️ BeautifulSoup not available — skipping website fetch")
            return ""
        try:
            if not url.startswith('http'):
                url = 'https://' + url
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
            }
            self.log(f"   Fetching website content from: {url}")
            response = requests.get(url, headers=headers, timeout=15)
            response.raise_for_status()
            soup = BeautifulSoup(response.text, 'html.parser')
            for script in soup(["script", "style", "nav", "footer", "header"]):
                script.decompose()
            text = soup.get_text(separator=' ', strip=True)
            text = re.sub(r'\s+', ' ', text).strip()
            if len(text) > 2000:
                text = text[:2000] + "..."
            self.log(f"   ✓ Fetched {len(text)} characters from website")
            return text
        except Exception as e:
            self.log(f"   ⚠️ Warning: Could not fetch website content: {e}")
            return ""

    def generate_background(self, client_name, website_link, title, keywords,
                            website_content, model="claude-sonnet-4-20250514") -> str:
        prompt = f"""Based on the following information, write a unique 200-word company/brand background that is SPECIFICALLY TAILORED to the blog post topic.
Client/Brand Name: {client_name}
Website: {website_link}
Blog Post Title: {title}
Target Keywords: {keywords}
Website Content (for reference):
{website_content if website_content else "No website content available - generate based on brand name and blog post context."}
CRITICAL Requirements:
- Write exactly around 200 words
- The background must be SPECIFICALLY RELEVANT to the blog post topic: "{title}"
- Focus on the company's expertise, products, or services that DIRECTLY RELATE to this specific topic
- Write in third person (e.g., "[Company Name] is...", "[Company Name] specializes in...")
- Include specific details about their capabilities related to: {keywords}
- Keep it professional and informative
- Do NOT include generic company overview - make it topic-focused
- Do NOT include any placeholder text or brackets
- Do NOT mention that this is AI-generated
Write only the background paragraph, nothing else."""
        try:
            self.log("   Generating 200-word topic-specific background...")
            message = self.client.messages.create(
                model=model,
                max_tokens=500,
                messages=[{"role": "user", "content": prompt}]
            )
            background = message.content[0].text.strip()
            self.log(f"   ✓ Background generated ({len(background.split())} words)")
            return background
        except Exception as e:
            self.log(f"   ✗ Error generating background: {e}")
            return f"{client_name} is a company providing products and services related to {title}."

    def generate_instruction(self, client_name, title, keywords, background,
                             website_content, model="claude-sonnet-4-20250514") -> str:
        prompt = f"""Based on the following information, write a detailed instruction/outline for a blog post writer.
Client/Brand Name: {client_name}
Blog Post Title: {title}
Target Keywords: {keywords}
Company Background: {background}
Website Content (for additional context):
{website_content if website_content else "No additional website content available."}
Write an instruction paragraph (50-100 words) that follows this EXACT pattern:
1. Start with "This blog post should focus on..." or "Focus this article on..."
2. Explain the CORE TOPIC the article should cover
3. Connect the topic to the client's products/services/expertise
4. Mention SPECIFIC ELEMENTS to highlight (features, benefits, comparisons, use cases, etc.)
5. Specify the TONE (educational, inspirational, authoritative, practical, etc.)
6. End with "The goal is to..." explaining the purpose for the reader
Requirements:
- Write 50-100 words
- Be specific to this exact blog post topic
- Do NOT use generic phrases - be specific to "{title}"
- Do NOT include any placeholder brackets
- Write as a single paragraph
Write only the instruction paragraph, nothing else."""
        try:
            self.log("   Generating instruction/outline...")
            message = self.client.messages.create(
                model=model,
                max_tokens=300,
                messages=[{"role": "user", "content": prompt}]
            )
            instruction = message.content[0].text.strip()
            self.log(f"   ✓ Instruction generated ({len(instruction.split())} words)")
            return instruction
        except Exception as e:
            self.log(f"   ✗ Error generating instruction: {e}")
            return (f"This blog post should focus on {title}, covering key aspects relevant to "
                    f"{client_name}'s expertise. The goal is to educate readers.")

    def create_dynamic_prompt(self, row_data: dict, custom_template: str = None) -> str:
        instruction = row_data.get('Instruction', '')
        template = custom_template if custom_template and custom_template.strip() else DEFAULT_ARTICLE_INSTRUCTIONS
        instructions_text = (
            template
            .replace('{KEYWORDS}', row_data.get('Keywords', ''))
            .replace('{CLIENT}', row_data.get('Client Name', ''))
            .replace('{LINKS}', row_data.get('Links To Add', ''))
        )
        return f"""Title: {row_data.get('Title', '')}
Keywords: {row_data.get('Keywords', '')}
Client: {row_data.get('Client Name', '')}
Company Background: {row_data.get('Background', '')}
Contact Link: {row_data.get('Links To Add', '')}
Website: {row_data.get('Website Link', '')}
Article Outline/Instructions:
{instruction if instruction else "No specific outline provided. Write a comprehensive article covering the topic thoroughly."}
{instructions_text}"""

    def send_to_anthropic(self, prompt: str, model: str = "claude-sonnet-4-20250514") -> str:
        message = self.client.messages.create(
            model=model,
            max_tokens=4000,
            messages=[{"role": "user", "content": prompt}]
        )
        return message.content[0].text

    def markdown_to_docx(self, content: str, title: str) -> "Document":
        doc = Document()
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        lines = content.split('\n')
        in_table = False
        table_rows = []

        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue

            if line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('|') and '|' in line:
                if not in_table:
                    in_table = True
                    table_rows = []
                row_data = [cell.strip() for cell in line.split('|') if cell.strip()]
                if row_data and not all('-' in cell for cell in row_data):
                    table_rows.append(row_data)
                next_is_table = (i + 1 < len(lines) and
                                 lines[i + 1].strip().startswith('|'))
                if not next_is_table and table_rows:
                    self.create_table(doc, table_rows)
                    in_table = False
                    table_rows = []
            elif re.match(r'^\d+\. ', line):
                para = doc.add_paragraph(style='List Number')
                self.add_formatted_text(para, re.sub(r'^\d+\. ', '', line))
            elif line.startswith(('- ', '* ', '• ')):
                para = doc.add_paragraph(style='List Bullet')
                self.add_formatted_text(para, line[2:])
            else:
                if not in_table:
                    para = doc.add_paragraph()
                    self.add_formatted_text(para, line)

            i += 1

        return doc

    def create_table(self, doc, table_rows):
        if not table_rows:
            return
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        table.style = 'Table Grid'
        for row_idx, row_data in enumerate(table_rows):
            for col_idx, cell_data in enumerate(row_data):
                if col_idx < len(table.rows[row_idx].cells):
                    cell = table.rows[row_idx].cells[col_idx]
                    clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_data)
                    cell.text = clean_text
                    if row_idx == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True

    def add_formatted_text(self, paragraph, text):
        text = re.sub(r'\*\*\[([^\]]+)\]\(([^)]+)\)\*\*', r'[[\1]](\2)', text)
        pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|\[\[[^\]]+\]\]\([^)]+\)|\[[^\]]+\]\([^)]+\))'
        parts = re.split(pattern, text)

        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif part.startswith('[[') and ']](' in part and part.endswith(')'):
                bracket_pos = part.find(']](')
                if bracket_pos > 0:
                    self.add_hyperlink(paragraph, part[bracket_pos + 3:-1], part[2:bracket_pos], bold=True)
            elif part.startswith('[') and '](' in part and part.endswith(')'):
                bracket_pos = part.find('](')
                if bracket_pos > 0:
                    self.add_hyperlink(paragraph, part[bracket_pos + 2:-1], part[1:bracket_pos])
            else:
                paragraph.add_run(part)

    def add_hyperlink(self, paragraph, url, text, bold=False):
        try:
            from docx.oxml.shared import qn
            from docx.oxml import parse_xml
            part = paragraph.part
            r_id = part.relate_to(
                url,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True
            )
            bold_xml = "<w:b/>" if bold else ""
            hyperlink_xml = f'''<w:hyperlink r:id="{r_id}" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:r><w:rPr>{bold_xml}<w:color w:val="0000FF"/><w:u w:val="single"/></w:rPr><w:t>{text}</w:t></w:r></w:hyperlink>'''
            paragraph._p.append(parse_xml(hyperlink_xml))
        except Exception:
            run = paragraph.add_run(text)
            run.bold = bold
            run.font.color.rgb = RGBColor(0, 0, 255)
            run.underline = True

    def save_article_as_docx(self, content: str, title: str, output_dir: str) -> str:
        safe_filename = self.sanitize_filename(title) or f"article_{int(time.time())}"
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)

        if DOCX_AVAILABLE:
            doc = self.markdown_to_docx(content, title)
            filepath = output_path / f"{safe_filename}.docx"
            doc.save(str(filepath))
            self.log(f"   Saved as DOCX: {filepath.name}")
            return str(filepath)
        else:
            filepath = output_path / f"{safe_filename}.txt"
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(f"{title}\n{'-' * len(title)}\n\n{content}")
            self.log(f"   Saved as TXT (install python-docx for .docx): {filepath.name}")
            return str(filepath)

    def process_csv_file(self, csv_file_path: str, output_dir: str,
                         delay_seconds: float = 3.0,
                         model: str = "claude-sonnet-4-20250514",
                         custom_template: str = None) -> list:
        """Process the CSV and generate articles, saving into output_dir."""
        try:
            try:
                df = pd.read_csv(csv_file_path, encoding='utf-8')
            except UnicodeDecodeError:
                df = pd.read_csv(csv_file_path, encoding='windows-1252')
            self.log(f"Loaded CSV with {len(df)} rows")
        except Exception as e:
            self.log(f"Error reading CSV file: {e}")
            return []

        required_cols = ['Client Name', 'Title', 'Keywords', 'Links To Add']
        optional_cols = ['Background', 'Website Link', 'Instruction', 'Status', 'Source Content']

        missing_cols = [col for col in required_cols if col not in df.columns]
        if missing_cols:
            self.log(f"Error: Missing required columns: {missing_cols}")
            return []

        for col in optional_cols:
            if col not in df.columns:
                df[col] = ''

        for col in optional_cols:
            df[col] = df[col].astype(str).replace('nan', '')

        df_clean = df.dropna(subset=['Client Name', 'Title'])
        df_clean = df_clean[df_clean['Title'].str.strip() != '']

        if 'Status' in df.columns:
            df_clean = df_clean[df_clean['Status'].str.upper().str.strip() == 'ACTIVE']

        if len(df_clean) == 0:
            self.log("⚠️ No rows with Status='ACTIVE' found. Nothing to process.")
            return []

        self.log(f"Processing {len(df_clean)} rows with Status='ACTIVE'")

        # Create client-specific folders inside output_dir
        unique_clients = df_clean['Client Name'].unique()
        client_folders = {}
        for client in unique_clients:
            folder_name = self.sanitize_folder_name(client)
            folder_path = Path(output_dir) / folder_name
            folder_path.mkdir(parents=True, exist_ok=True)
            client_folders[client] = folder_path
            self.log(f"📁 Created folder: {folder_name}")

        results = []

        for idx, (index, row) in enumerate(df_clean.iterrows()):
            row_dict = row.to_dict()
            title = row_dict['Title']
            client = row_dict['Client Name']
            keywords = row_dict.get('Keywords', '')
            website_link = row_dict.get('Website Link', '')
            client_folder = client_folders.get(client, Path(output_dir))

            self.log(f"\n{'='*50}")
            self.log(f"Article {idx + 1}/{len(df_clean)}: {title}")
            self.log(f"Client: {client}")
            self.log(f"{'='*50}")

            try:
                # Step 1: Get source content
                source_content = row_dict.get('Source Content', '')
                website_content = ""

                if source_content and str(source_content).strip() and str(source_content).lower() != 'nan':
                    website_content = str(source_content).strip()[:2000]
                    self.log(f"   ✓ Using provided Source Content ({len(website_content)} chars)")
                elif website_link and str(website_link).strip() and str(website_link).lower() != 'nan':
                    self.log("   📡 Fetching from website...")
                    website_content = self.fetch_website_content(str(website_link).strip())
                    time.sleep(delay_seconds)
                else:
                    self.log("   ⚠️ No Source Content or website link — generating from title/keywords only")

                # Step 2: Generate background
                self.log("\n📝 Step 1/3: Generating background...")
                background = self.generate_background(
                    client_name=client, website_link=str(website_link) if website_link else "",
                    title=title, keywords=keywords, website_content=website_content, model=model
                )
                row_dict['Background'] = background
                df.at[index, 'Background'] = background
                time.sleep(delay_seconds)

                # Step 3: Generate instruction
                self.log("\n📝 Step 2/3: Generating instruction...")
                instruction = self.generate_instruction(
                    client_name=client, title=title, keywords=keywords,
                    background=background, website_content=website_content, model=model
                )
                row_dict['Instruction'] = instruction
                df.at[index, 'Instruction'] = instruction
                time.sleep(delay_seconds)

                # Step 4: Generate article
                self.log("\n📝 Step 3/3: Generating article...")
                prompt = self.create_dynamic_prompt(row_dict, custom_template=custom_template)
                article_content = self.send_to_anthropic(prompt, model)

                # Step 5: Save
                saved_path = self.save_article_as_docx(article_content, title, str(client_folder))
                self.log(f"✅ Article saved: {Path(saved_path).name}")

                results.append({
                    'index': index, 'client': client, 'title': title,
                    'saved_path': saved_path, 'success': True, 'error': None
                })

            except Exception as e:
                self.log(f"❌ Error: {e}")
                results.append({
                    'index': index, 'client': client, 'title': title,
                    'saved_path': None, 'success': False, 'error': str(e)
                })

            if idx < len(df_clean) - 1:
                self.log(f"\n⏳ Waiting {delay_seconds}s before next article...")
                time.sleep(delay_seconds)

        # Save updated CSV to output dir
        try:
            updated_csv = Path(output_dir) / "articles_updated.csv"
            df.to_csv(str(updated_csv), index=False, encoding='utf-8-sig')
            self.log(f"\n📄 Updated CSV saved: {updated_csv.name}")
        except Exception as e:
            self.log(f"\n⚠️ Could not save updated CSV: {e}")

        return results
